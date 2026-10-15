"""Images in documents, described by a vision model on the user's machine.

Text extraction misses what a drawing or photo shows. When a vision model is
set (Instellingen → Indexeren), the NAS marks the places that need a
description - PDF pages that are mostly image or drawing, and the larger
images in Word documents - with a placeholder in the extracted text. The
indexer (browser or CLI) fetches each one as a PNG (GET /api/index/visual),
has its Ollama describe it, and sends the descriptions back with the next
prepare. They take the placeholders' place, so they are split, embedded and
found like any other text, under the right heading and page.

No model runs here: the NAS only renders pages (pypdfium2, through
pdfplumber) and reads embedded images (Pillow).
"""

import base64
import io
import re
from pathlib import Path

# Longest side of an image sent to the vision model: enough to read a title
# block, small enough to describe in seconds rather than minutes.
MAX_SIDE = 1536
# Word images below this many bytes are logos and icons, not content.
MIN_DOCX_IMAGE_BYTES = 20_000

PROMPT = (
    "Beschrijf deze afbeelding uit een document, voor een zoekindex. Antwoord in het Nederlands. "
    "Noem eerst wat voor afbeelding het is (bijvoorbeeld bouwtekening, plattegrond, gevel, doorsnede, "
    "detail, schema, grafiek, foto of scan). Neem daarna over wat er in een titelblok of legenda staat "
    "(project, tekeningnummer, revisie, schaal, datum, opdrachtgever) en alle leesbare tekst en labels, "
    "zoals ruimtenamen. Beschrijf tot slot in een paar zinnen wat er te zien is. Wees feitelijk: "
    "noem alleen maten en getallen die je duidelijk kunt lezen en verzin niets."
)

_PLACEHOLDER = re.compile(r"\[\[afbeelding:([a-z]\d+)\]\]")


def placeholder(visual_id: str) -> str:
    return f"[[afbeelding:{visual_id}]]"


def label(visual_id: str) -> str:
    """"p3" -> "pagina 3", "w2" -> "afbeelding 2"."""
    kind, number = visual_id[0], visual_id[1:]
    return f"pagina {number}" if kind == "p" else f"afbeelding {number}"


def ids(text: str) -> list[str]:
    return _PLACEHOLDER.findall(text)


def fill(text: str, descriptions: dict[str, str] | None, keep: set[str] | None = None) -> str:
    """Put each description in place of its placeholder; drop placeholders without one."""
    descriptions = descriptions or {}

    def replace(match: re.Match) -> str:
        visual_id = match.group(1)
        description = (descriptions.get(visual_id) or "").strip()
        if not description or (keep is not None and visual_id not in keep):
            return ""
        return f"\n\nBeschrijving van {label(visual_id)}:\n{description}\n\n"

    return re.sub(r"\n{3,}", "\n\n", _PLACEHOLDER.sub(replace, text)).strip()


def pdf_page_needs_description(page) -> bool:
    """A page that is mostly image, a (vector) drawing, or graphics with hardly any text."""
    area = float(page.width * page.height) or 1.0
    image_area = sum(max(0, (i["x1"] - i["x0"]) * (i["bottom"] - i["top"])) for i in page.images)
    if image_area >= 0.2 * area:
        return True
    graphics = len(page.curves) + len(page.lines) + len(page.rects)
    # Drawings are thousands of strokes; a ruled table is a few hundred.
    if len(page.curves) >= 50 or graphics >= 500:
        return True
    return len(page.chars) < 100 and (bool(page.images) or graphics >= 20)


def paragraph_images(document, paragraph) -> list[bytes]:
    """The content images (not logos or icons) in one Word paragraph."""
    blobs = []
    for rel_id in paragraph._p.xpath(".//a:blip/@r:embed"):
        blob = getattr(document.part.related_parts.get(rel_id), "blob", b"")
        if len(blob) >= MIN_DOCX_IMAGE_BYTES:
            blobs.append(blob)
    return blobs


def docx_images(document):
    """Every content image in the body paragraphs, in document order: "w1", "w2", ..."""
    for block in document.iter_inner_content():
        if hasattr(block, "_p"):
            yield from paragraph_images(document, block)


def _png(image) -> str:
    image = image.convert("RGB")
    image.thumbnail((MAX_SIDE, MAX_SIDE))
    out = io.BytesIO()
    image.save(out, format="PNG", optimize=True)
    return base64.b64encode(out.getvalue()).decode("ascii")


def render(path: Path, visual_id: str) -> str:
    """The image for a placeholder, as base64 PNG; KeyError when there is none."""
    kind, number = visual_id[0], int(visual_id[1:])
    ext = path.suffix.lower()
    if kind == "p" and ext == ".pdf":
        import pdfplumber

        with pdfplumber.open(str(path)) as pdf:
            if not 1 <= number <= len(pdf.pages):
                raise KeyError(visual_id)
            page = pdf.pages[number - 1]
            resolution = 72 * MAX_SIDE / max(float(page.width), float(page.height))
            return _png(page.to_image(resolution=resolution).original)
    if kind == "w" and ext in (".docx", ".doc"):
        import docx
        from PIL import Image

        for i, blob in enumerate(docx_images(docx.Document(str(path))), start=1):
            if i == number:
                return _png(Image.open(io.BytesIO(blob)))
    raise KeyError(visual_id)
