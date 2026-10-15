"""Tables in documents: read as Markdown tables, kept whole with their header when split."""

from pathlib import Path

import docx
import openpyxl

from app import documents, tables

FIXTURES = Path(__file__).parent / "fixtures"
CHUNKS = {"parent_chunk_size": 1500, "child_chunk_size": 200, "parent_chunk_overlap": 300}


def test_markdown_table_cleans_cells():
    rows = [["Naam", None, "Bedrag"], ["a|b", None, "1\n200"], [None, None, None]]
    assert tables.markdown_table(rows) == "| Naam | Bedrag |\n|---|---|\n| a\\|b | 1 200 |"


def test_large_tables_are_cut_into_pieces_with_the_header():
    rows = [f"| 2024-{i:04d} | Klant {i} | € {i * 100},00 |" for i in range(60)]
    text = "Intro.\n\n| Factuur | Klant | Bedrag |\n|---|---|---|\n" + "\n".join(rows) + "\n\nSlot."
    pieces = [p for p in tables.keep_tables_whole(text, 100).split(tables.PARAGRAPH) if "| Factuur |" in p]
    assert len(pieces) > 2
    assert all(p.startswith("| Factuur | Klant | Bedrag |\n|---|---|---|\n| 2024-") for p in pieces)
    assert sum(p.count("| 2024-") for p in pieces) == 60


def test_pdf_tables_stay_tables_and_every_chunk_keeps_the_header():
    text = documents.extract_text(FIXTURES / "facturen.pdf")
    assert "Hieronder staan alle facturen" in text
    assert "| Factuur | Klant | Bedrag | Vervaldatum |" in text
    assert "| 2024-0012 | Klant 1 | € 1200,00 | 12 maart |" in text
    assert text.index("Hieronder staan") < text.index("| Factuur |") < text.index("Betaling binnen 30 dagen")

    leaves = [n for n in documents.split(text, {"file_name": "facturen.pdf"}, "h", CHUNKS) if n["is_leaf"]]
    with_rows = [n["content"] for n in leaves if "| 2024-" in n["content"]]
    assert len(with_rows) > 1
    assert all("| Factuur | Klant | Bedrag | Vervaldatum |" in c for c in with_rows)
    # No row is cut in half between chunks.
    for content in with_rows:
        assert all(line.endswith("|") for line in content.splitlines() if line.startswith("| 2024-"))


def test_word_tables_in_document_order(tmp_path):
    d = docx.Document()
    d.add_heading("Prijslijst", 1)
    d.add_paragraph("Tarieven per uur.")
    t = d.add_table(rows=2, cols=2)
    for r, (a, b) in enumerate([("Dienst", "Tarief"), ("Advies", "€ 95")]):
        t.cell(r, 0).text, t.cell(r, 1).text = a, b
    d.add_paragraph("Prijzen exclusief btw.")
    d.save(tmp_path / "p.docx")
    text = documents.extract_text(tmp_path / "p.docx")
    assert text.index("# Prijslijst") < text.index("| Dienst | Tarief |") < text.index("| Advies | € 95 |") < text.index("exclusief btw")


def test_excel_sheets_and_csv(tmp_path):
    wb = openpyxl.Workbook()
    wb.active.title = "Uren"
    for row in [("Datum", "Project", "Uren"), ("2024-03-01", "JEX", 6)]:
        wb.active.append(row)
    wb.save(tmp_path / "u.xlsx")
    assert documents.extract_text(tmp_path / "u.xlsx") == "# Uren\n\n| Datum | Project | Uren |\n|---|---|---|\n| 2024-03-01 | JEX | 6 |"

    (tmp_path / "k.csv").write_text("Post;Bedrag\nHuur;1200\n", encoding="utf-8")
    assert documents.extract_text(tmp_path / "k.csv") == "| Post | Bedrag |\n|---|---|\n| Huur | 1200 |"
